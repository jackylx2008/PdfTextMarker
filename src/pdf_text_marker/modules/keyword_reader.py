"""读取 XLSX 或 CSV 中的关键词。"""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from pdf_text_marker.models import ExcelKeywordSource, KeywordSpec


KEYWORD_COLUMN_CANDIDATES = ("keyword", "关键词")
PDF_COLUMN_CANDIDATES = ("pdf_name", "pdf文件", "pdf文件名", "PDF文件", "PDF文件名")


class KeywordFileError(ValueError):
    """关键词文件格式或内容无效。"""


@dataclass(frozen=True, slots=True)
class KeywordSourceInfo:
    """关键词文件可供界面选择的结构信息。"""

    sheets: tuple[str, ...]
    selected_sheet: str | None
    columns: tuple[str, ...]


def inspect_keyword_source(path: Path, sheet_name: str | None = None) -> KeywordSourceInfo:
    """返回工作表和表头，不读取业务数据。"""
    path = path.resolve()
    _validate_source(path)
    if path.suffix.casefold() == ".xlsx":
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            sheets = tuple(workbook.sheetnames)
            selected = _select_sheet(sheets, sheet_name)
            worksheet = workbook[selected]
            first_row = next(worksheet.iter_rows(values_only=True), ())
            columns = _headers(first_row)
        finally:
            workbook.close()
        return KeywordSourceInfo(sheets=sheets, selected_sheet=selected, columns=columns)

    rows = _read_csv_rows(path)
    return KeywordSourceInfo(sheets=(), selected_sheet=None, columns=_headers(rows[0] if rows else ()))


def read_keywords(
    path: Path,
    sheet_name: str | None = None,
    keyword_column: str | None = None,
    pdf_name_column: str | None = None,
) -> list[KeywordSpec]:
    """读取、清理并去重关键词。"""
    path = path.resolve()
    _validate_source(path)
    is_xlsx = path.suffix.casefold() == ".xlsx"
    if is_xlsx:
        headers, rows = _read_xlsx_rows(path, sheet_name)
    else:
        matrix = _read_csv_rows(path)
        if not matrix:
            raise KeywordFileError(f"关键词文件为空: {path}")
        headers, rows = _headers(matrix[0]), matrix[1:]

    if is_xlsx:
        keyword_index = _find_excel_column(headers, keyword_column, KEYWORD_COLUMN_CANDIDATES, required=True)
        pdf_index = _find_excel_column(headers, pdf_name_column, PDF_COLUMN_CANDIDATES, required=False)
    else:
        keyword_index = _find_column(headers, keyword_column, KEYWORD_COLUMN_CANDIDATES, required=True)
        pdf_index = _find_column(headers, pdf_name_column, PDF_COLUMN_CANDIDATES, required=False)
    keywords: list[KeywordSpec] = []
    seen: set[tuple[str, str]] = set()
    for row in rows:
        keyword = _keyword_text(row[keyword_index] if keyword_index < len(row) else None)
        if not keyword:
            continue
        target = _cell_text(row[pdf_index] if pdf_index is not None and pdf_index < len(row) else None)
        item = KeywordSpec(keyword=keyword, pdf_name=target or None)
        if item.identity not in seen:
            seen.add(item.identity)
            keywords.append(item)
    if not keywords:
        raise KeywordFileError("关键词列中没有有效内容")
    return keywords


def read_excel_keywords(
    path: Path,
    sources: Iterable[ExcelKeywordSource],
    default_pdf_name_column: str | None = None,
) -> list[KeywordSpec]:
    """合并多个 Excel Sheet/列中的关键词并统一去重。"""
    path = path.resolve()
    if path.suffix.casefold() != ".xlsx":
        raise KeywordFileError("excel_sources 只能用于 .xlsx 文件")
    combined: list[KeywordSpec] = []
    seen: set[tuple[str, str]] = set()
    source_count = 0
    for source in sources:
        source_count += 1
        for column in source.keyword_columns:
            items = read_keywords(
                path,
                sheet_name=source.sheet_name,
                keyword_column=column,
                pdf_name_column=source.pdf_name_column or default_pdf_name_column,
            )
            for item in items:
                if item.identity not in seen:
                    seen.add(item.identity)
                    combined.append(item)
    if source_count == 0:
        raise KeywordFileError("excel_sources 至少需要一个工作表配置")
    if not combined:
        raise KeywordFileError("配置的 Excel 工作表和列中没有有效关键词")
    return combined


def _read_xlsx_rows(path: Path, sheet_name: str | None) -> tuple[tuple[str, ...], list[tuple[object, ...]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        selected = _select_sheet(tuple(workbook.sheetnames), sheet_name)
        matrix = list(workbook[selected].iter_rows(values_only=True))
    finally:
        workbook.close()
    if not matrix:
        raise KeywordFileError(f"工作表为空: {sheet_name or '默认工作表'}")
    return _headers(matrix[0]), matrix[1:]


def _read_csv_rows(path: Path) -> list[list[str]]:
    text: str | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise KeywordFileError(f"CSV 编码无法识别（支持 UTF-8 和 GB18030）: {path}")
    try:
        dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [row for row in csv.reader(text.splitlines(), dialect) if any(cell.strip() for cell in row)]


def _validate_source(path: Path) -> None:
    if not path.is_file():
        raise KeywordFileError(f"关键词文件不存在: {path}")
    if path.suffix.casefold() not in {".xlsx", ".csv"}:
        raise KeywordFileError("关键词文件只支持 .xlsx 和 .csv")


def _select_sheet(sheets: tuple[str, ...], requested: str | None) -> str:
    if not sheets:
        raise KeywordFileError("Excel 文件不包含工作表")
    if not requested:
        return sheets[0]
    for sheet in sheets:
        if sheet.casefold() == requested.casefold():
            return sheet
    raise KeywordFileError(f"找不到工作表 {requested}；可选值: {', '.join(sheets)}")


def _headers(values: Iterable[object]) -> tuple[str, ...]:
    headers = tuple(_cell_text(value) for value in values)
    if not headers or not any(headers):
        raise KeywordFileError("首行必须包含列名")
    return headers


def _find_column(
    headers: tuple[str, ...],
    requested: str | None,
    candidates: tuple[str, ...],
    *,
    required: bool,
) -> int | None:
    if requested:
        matches = [index for index, header in enumerate(headers) if header.casefold() == requested.strip().casefold()]
        if len(matches) > 1:
            raise KeywordFileError(f"列名 {requested} 重复，请改用 Excel 列字母指定")
        if not matches:
            raise KeywordFileError(f"找不到列 {requested}；可选值: {', '.join(filter(None, headers))}")
        return matches[0]
    for name in candidates:
        matches = [index for index, header in enumerate(headers) if header.casefold() == name.casefold()]
        if len(matches) > 1:
            raise KeywordFileError(f"自动识别到多个 {name} 列，请明确指定列")
        if matches:
            return matches[0]
    if required:
        raise KeywordFileError(
            "找不到关键词列。请使用列名 keyword/关键词，或在界面和参数中明确选择关键词列"
        )
    return None


def _find_excel_column(
    headers: tuple[str, ...],
    requested: str | None,
    candidates: tuple[str, ...],
    *,
    required: bool,
) -> int | None:
    """Excel 列既可按首行列名指定，也可直接使用 A、B、AA 等列字母。"""
    if requested:
        requested_text = requested.strip()
        for index, header in enumerate(headers):
            if header and header.casefold() == requested_text.casefold():
                return index
        normalized = requested_text.upper().replace("$", "")
        if normalized.endswith(":"):
            normalized = normalized[:-1]
        if ":" in normalized:
            left, right = normalized.split(":", 1)
            if left == right:
                normalized = left
        try:
            index = column_index_from_string(normalized) - 1 if normalized.isalpha() and len(normalized) <= 3 else -1
        except ValueError:
            index = -1
        if index >= 0:
            if index >= len(headers):
                raise KeywordFileError(f"Excel 列 {requested} 超出工作表范围")
            return index
    return _find_column(headers, requested, candidates, required=required)


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _keyword_text(value: object) -> str:
    """清理关键词，并删除首个中英文冒号及其后的说明。"""
    text = _cell_text(value)
    return re.split(r"[:：]", text, maxsplit=1)[0].strip()
